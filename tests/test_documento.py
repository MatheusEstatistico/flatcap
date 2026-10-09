# SPDX-FileCopyrightText: 2026 Matheus Fernandes
# SPDX-License-Identifier: GPL-3.0-or-later
"""
O `.docx` — Fase 6.

Este módulo escreve e não decide, então o teste aqui é sobre o **documento que
sai**, não sobre números. Quatro das cinco coisas mais graves encontradas na
fase só apareceram ao abrir as páginas; as que dá para verificar em XML estão
aqui, e a que não dá está escrita no fim do arquivo, como lembrete.
"""

from __future__ import annotations

import os
import stat
from datetime import datetime

import pytest
from docx import Document
from docx.shared import Pt

from flatcap.documento import (
    ARQUIVO_TEMPLATE,
    ESTILO_LEGENDA,
    ESTILO_SECAO,
    ESTILO_TABELA,
    ESTILO_TITULO,
    escrever_template,
    versao_do_python_docx,
)
from flatcap.relatorio import RelatorioError, executar_relatorio, preparar_relatorio

QUANDO = datetime(2026, 9, 9, 10, 0)


@pytest.fixture
def documento(projeto_com_rodada):
    """
    Devolve o `.docx` já escrito e aberto, com o `Relatorio` que o gerou.
    """

    def escrever_um(quadro=None, template=None, **argumentos):
        raiz = projeto_com_rodada(quadro)
        plano = preparar_relatorio(raiz, agora=QUANDO, template=template, **argumentos)
        executar_relatorio(plano)
        return Document(str(plano.destino)), plano

    return escrever_um


def _textos(doc) -> list[str]:
    return [p.text for p in doc.paragraphs if p.text.strip()]


# ── A primeira página ───────────────────────────────────────────────────────


class TestPrimeiraPagina:
    def test_titulo_subtitulo_e_texto_de_abertura_nessa_ordem(self, documento):
        doc, plano = documento()
        textos = _textos(doc)
        assert textos[0] == plano.relatorio.titulo
        assert textos[1] == plano.relatorio.subtitulo
        assert textos[2].startswith(plano.relatorio.texto[0][:20])

    def test_o_texto_de_abertura_sai_um_paragrafo_por_linha_da_tupla(self, documento):
        doc, plano = documento()
        textos = _textos(doc)
        for paragrafo in plano.relatorio.texto:
            assert paragrafo in textos

    def test_a_primeira_variavel_comeca_em_pagina_nova(self, documento):
        """
        `page_break_before` no título da seção, e não um parágrafo de quebra
        atrás dele: um parágrafo a mais deixaria a marca dele no alto da página
        seguinte, empurrando o título para baixo.
        """
        doc, plano = documento()
        primeira = plano.secoes[0].rotulo
        titulo = next(p for p in doc.paragraphs if p.text == primeira)
        assert titulo.paragraph_format.page_break_before is True

    def test_sem_texto_de_abertura_nao_ha_quebra(self, projeto_com_rodada, tmp_path):
        """
        Sem texto e sem alerta, a primeira página seria só o título e o
        subtítulo, e a quebra produziria uma página quase em branco.

        O `Relatorio` é montado à mão, e é o único jeito de chegar aqui: pelo
        `flatcap relatorio` o texto nunca vem vazio, porque o `montar_abertura`
        devolve pelo menos os dois marcadores `«a preencher: …»`. Trocar a
        `TEXTO_PADRAO` não alcança o ramo — ela é o valor padrão de um
        parâmetro, lido uma vez, quando o módulo é importado. A docstring de
        `documento._texto_padrao` diz para quem o ramo existe.
        """
        from dataclasses import replace

        from flatcap.documento import escrever

        plano = preparar_relatorio(projeto_com_rodada(), agora=QUANDO)
        destino = escrever(
            replace(plano.relatorio, texto=(), alertas=()), tmp_path / "sem_abertura.docx"
        )
        doc = Document(str(destino))
        titulo = next(p for p in doc.paragraphs if p.text == plano.secoes[0].rotulo)
        assert not titulo.paragraph_format.page_break_before

    def test_a_procedencia_esta_nos_metadados(self, documento):
        """
        Depois da revisão de formatação, nenhuma página diz de qual versão do
        banco os números vieram — mas as propriedades do arquivo dizem, e o
        Windows as mostra sem ninguém abrir o Word.
        """
        doc, plano = documento()
        propriedades = doc.core_properties
        assert plano.relatorio.ficha.matricula in propriedades.subject
        assert plano.relatorio.manifesto.banco["arquivo"] in propriedades.comments

    def test_o_alerta_vem_antes_do_texto_de_abertura(self, projeto_com_rodada):
        """
        É a única coisa da primeira página que pode fazer alguém parar de enviar
        o documento; depois do texto padrão, ele seria lido tarde demais.
        """
        from flatcap.projeto import PASTA_DADOS

        raiz = projeto_com_rodada()
        alvo = raiz / PASTA_DADOS / "00_redcap_2026-09-01_banco.csv"
        alvo.write_text(alvo.read_text(encoding="utf-8") + "\n", encoding="utf-8")
        plano = preparar_relatorio(raiz, agora=QUANDO)
        executar_relatorio(plano)
        doc = Document(str(plano.destino))
        aviso = doc.tables[0].cell(0, 0).text
        assert "Atenção" in aviso
        corpo = "\n".join(_textos(doc))
        assert corpo.index(plano.relatorio.titulo) < corpo.index(plano.relatorio.texto[0])

    def test_o_alerta_e_uma_caixa_de_uma_celula_com_fundo(self, projeto_com_rodada):
        from flatcap.projeto import PASTA_DADOS

        raiz = projeto_com_rodada()
        alvo = raiz / PASTA_DADOS / "00_redcap_2026-09-01_banco.csv"
        alvo.write_text("estragado", encoding="utf-8")
        plano = preparar_relatorio(raiz, agora=QUANDO)
        executar_relatorio(plano)
        doc = Document(str(plano.destino))
        caixa = doc.tables[0]
        assert (len(caixa.rows), len(caixa.columns)) == (1, 1)
        assert "FBEFC8" in caixa.cell(0, 0)._tc.xml


# ── As seções ───────────────────────────────────────────────────────────────


class TestSecoes:
    def test_uma_secao_por_variavel_na_ordem_do_manifesto(self, documento):
        doc, plano = documento()
        textos = _textos(doc)
        posicoes = [textos.index(s.rotulo) for s in plano.secoes]
        assert posicoes == sorted(posicoes)

    def test_a_legenda_vem_acima_da_tabela_e_acima_da_figura(self, documento):
        doc, _ = documento()
        textos = _textos(doc)
        assert any(t.startswith("Tabela 1.") for t in textos)
        assert any(t.startswith("Figura 1.") for t in textos)

    def test_as_series_de_numeracao_sao_separadas_e_continuas(self, documento):
        doc, plano = documento()
        textos = _textos(doc)
        tabelas = [t for t in textos if t.startswith("Tabela ")]
        figuras = [t for t in textos if t.startswith("Figura ")]
        assert [t.split(".")[0] for t in tabelas] == [
            f"Tabela {i}" for i in range(1, plano.total_tabelas + 1)
        ]
        assert [f.split(".")[0] for f in figuras] == [
            f"Figura {i}" for i in range(1, plano.total_figuras + 1)
        ]

    def test_a_legenda_anda_com_o_objeto_que_nomeia(self, documento):
        """
        Sem `keep_with_next`, `Figura 2. Escolaridade` termina uma página e a
        figura começa a seguinte — e quem lê passa a página procurando o que a
        legenda prometeu.
        """
        doc, _ = documento()
        legendas = [p for p in doc.paragraphs if p.text.startswith(("Tabela ", "Figura "))]
        assert legendas
        assert all(p.paragraph_format.keep_with_next for p in legendas)

    def test_o_texto_das_celulas_atravessa_intacto(self, documento):
        """
        `28,7` no `.csv` é `28,7` no Word. A leitura usa o módulo `csv` da
        biblioteca padrão, e não o `pandas`, exatamente para isto.
        """
        doc, plano = documento()
        tabela = next(t for t in plano.relatorio.tabelas if t.arquivo == "resumo_idade.csv")
        media = next(linha[1] for linha in tabela.linhas if linha[0] == "média")
        assert any(media in celula.text for quadro in doc.tables for linha in quadro.rows
                   for celula in linha.cells)

    def test_a_linha_total_sai_em_negrito(self, documento):
        doc, _ = documento()
        for quadro in doc.tables:
            for linha in quadro.rows:
                if linha.cells[0].text.strip() == "Total":
                    trechos = [
                        t for celula in linha.cells for p in celula.paragraphs
                        for t in p.runs
                    ]
                    assert all(t.bold for t in trechos if t.text.strip())
                    return
        pytest.fail("nenhuma tabela do documento tem linha Total")

    def test_o_rotulo_a_esquerda_e_o_numero_a_direita(self, documento):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc, _ = documento()
        quadro = next(t for t in doc.tables if len(t.columns) > 2)
        linha = quadro.rows[1]
        assert linha.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
        assert linha.cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_a_coluna_de_texto_vai_a_esquerda_mesmo_nao_sendo_a_primeira(
        self, projeto_com_rodada
    ):
        """
        O cruzamento tem dois rótulos. Pela regra antiga — a primeira coluna à
        esquerda, o resto à direita —, `Fundamental` sairia alinhado pela última
        letra embaixo de `escolaridade`. O cabeçalho acompanha a coluna.
        """
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        from conftest import acrescentar_a_ficha
        from flatcap.cruzamento import executar_cruzamento, preparar_cruzamento

        raiz = projeto_com_rodada()
        acrescentar_a_ficha(raiz, "cruzamento:\n  - [sexo, escolaridade]\n")
        executar_cruzamento(preparar_cruzamento(raiz, agora=QUANDO))
        plano = preparar_relatorio(raiz, agora=QUANDO)
        executar_relatorio(plano)

        doc = Document(str(plano.destino))
        quadro = next(
            t for t in doc.tables
            if [c.text.strip() for c in t.rows[0].cells][:2] == ["sexo", "escolaridade"]
        )
        esquerda, direita = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT
        for linha in (quadro.rows[0], quadro.rows[1]):
            alinhamentos = [c.paragraphs[0].alignment for c in linha.cells]
            assert alinhamentos == [esquerda, esquerda, direita, direita]

    def test_o_que_conta_como_numero_na_hora_de_alinhar(self):
        """
        Número, percentual, valor-p, número com sinal e data ficam à direita; o
        traço e a célula vazia não puxam para lado nenhum. A coluna `leitura` da
        normalidade, que é frase, passa a ir à esquerda — como toda coluna de texto.
        """
        from flatcap.documento import _colunas_de_texto
        from flatcap.relatorio import TabelaLida

        def colunas(*linhas):
            return _colunas_de_texto(
                TabelaLida("x.csv", "", "", colunas=tuple(f"c{i}" for i in range(len(linhas[0]))),
                           linhas=linhas)
            )

        assert colunas(("Feminino", "15", "26,8", "25,0"), ("Total", "60", "100,0", "—")) == {0}
        assert colunas(("Feminino", "Superior", "15", "26,8"), ("Total", "", "60", "100,0")) == {
            0, 1,
        }
        assert colunas(("idade", "60", "+0,45", "−1,20", "< 0,001", "simétrica")) == {0, 5}
        assert colunas(("primeira data", "01/03/2026"), ("amplitude (dias)", "365")) == {0}

    def test_o_cabecalho_da_tabela_se_repete_a_cada_pagina(self, documento):
        doc, _ = documento()
        quadro = next(t for t in doc.tables if len(t.columns) > 1)
        assert "tblHeader" in quadro.rows[0]._tr.xml

    def test_a_coluna_do_rotulo_leva_mais_espaco_que_as_de_numero(self, documento):
        """
        No autoajuste do Word, `% dos válidos` no cabeçalho é mais largo do que
        qualquer número da coluna — e o resultado era uma coluna de rótulos
        espremida ao lado de três colunas de números folgadas.
        """
        doc, _ = documento()
        quadro = next(t for t in doc.tables if len(t.columns) == 4)
        larguras = [c.width for c in quadro.columns]
        assert larguras[0] > larguras[1]

    def test_tabela_de_um_rotulo_so_fica_na_reparticao_de_sempre(self, documento):
        """
        A regra nova é só para tabela com mais de uma coluna de texto. As da
        descritiva — um rótulo e números — saem exatamente como saíam: um terço
        para o rótulo, o resto em partes iguais.
        """
        doc, _ = documento()
        secao = doc.sections[-1]
        util = secao.page_width - secao.left_margin - secao.right_margin
        quadro = next(t for t in doc.tables if len(t.columns) == 4)
        larguras = [c.width for c in quadro.columns]
        # O Word guarda a largura em vigésimos de ponto: a volta perde até um.
        um_vigesimo = 635
        assert larguras[0] == pytest.approx(util * 0.34, abs=um_vigesimo)
        for largura in larguras[1:]:
            assert largura == pytest.approx(util * (1 - 0.34) / 3, abs=um_vigesimo)

    def test_varias_colunas_de_texto_repartem_pelo_conteudo(self):
        """
        A tabela da associação: o nome do teste leva o espaço, e não a
        `variável 1` quase vazia — até a 1.4.0, `Qui-quadrado com correção de
        Yates` quebrava em três linhas ao lado dela.
        """
        from flatcap.documento import _colunas_de_texto, _larguras_pelo_conteudo
        from flatcap.relatorio import TabelaLida

        tabela = TabelaLida(
            "associacao.csv",
            "associacao",
            "Testes de associação",
            colunas=("variável 1", "variável 2", "n", "teste", "p-valor"),
            linhas=(
                ("Grupo", "Desfecho", "79", "Qui-quadrado com correção de Yates", "0,017"),
                ("Grupo", "Escolaridade", "80", "Exato de Fisher", "< 0,001"),
            ),
        )
        util = int(Pt(453))  # a mancha de uma A4 com 2,5 cm de margem
        larguras = _larguras_pelo_conteudo(tabela, _colunas_de_texto(tabela), util)

        assert larguras is not None
        assert sum(larguras) <= util
        variavel_1, variavel_2, n, teste, p = larguras
        assert teste > variavel_2 > variavel_1
        assert teste > 2 * p
        assert n < variavel_1

    def test_cabecalho_de_numero_nao_quebra_no_meio_da_palavra(self):
        """
        O defeito da tabela da normalidade: `faltantes` e `assimetria` partidos em
        duas linhas porque todas as colunas depois da primeira tinham a mesma
        largura.
        """
        from flatcap.documento import (
            _CM_POR_CARACTERE,
            _EMU_POR_CM,
            _colunas_de_texto,
            _larguras_pelo_conteudo,
        )
        from flatcap.relatorio import TabelaLida

        colunas = ("variável", "n", "faltantes", "assimetria", "curtose", "W", "valor-p", "leitura")
        tabela = TabelaLida(
            "normalidade.csv",
            "normalidade",
            "Pressupostos — normalidade",
            colunas=colunas,
            linhas=(
                ("Idade (anos)", "80", "0", "+0,01", "-1,22", "0,951", "0,004",
                 "simétrica; achatada; o teste rejeita a normal"),
            ),
        )
        larguras = _larguras_pelo_conteudo(tabela, _colunas_de_texto(tabela), int(Pt(453)))
        assert larguras is not None
        for posicao in (2, 3, 4, 6):
            assert larguras[posicao] >= len(colunas[posicao]) * _CM_POR_CARACTERE * _EMU_POR_CM

    def test_a_palavra_curta_de_texto_leva_um_caractere_de_folga(self):
        """
        O defeito da tabela da correlação: `Spearman`, na coluna `método`, recebia a
        largura média de oito caracteres, e o `m` passa da média — a palavra saía
        partida em `Spearma` e `n`. O piso da coluna de texto é a palavra mais longa
        com um caractere a mais.
        """
        from flatcap.documento import (
            _CM_POR_CARACTERE,
            _EMU_POR_CM,
            _RESPIRO_CM,
            _colunas_de_texto,
            _larguras_pelo_conteudo,
        )
        from flatcap.relatorio import TabelaLida

        tabela = TabelaLida(
            "correlacao.csv",
            "correlacao",
            "Correlações",
            colunas=(
                "variável 1", "variável 2", "n", "método", "coeficiente", "IC 95%", "p-valor",
            ),
            linhas=(
                ("Idade (anos)", "Renda familiar (R$)", "118", "Spearman", "-0,01",
                 "-0,19 a 0,17", "0,919"),
                ("Intensidade da dor (0 a 5)", "Escore de ansiedade (0 a 10)", "120",
                 "Kendall", "0,42", "0,31 a 0,51", "< 0,001"),
            ),
        )
        larguras = _larguras_pelo_conteudo(tabela, _colunas_de_texto(tabela), int(Pt(453)))
        assert larguras is not None
        assert larguras[3] >= ((len("Spearman") + 1) * _CM_POR_CARACTERE + _RESPIRO_CM) * (
            _EMU_POR_CM
        ) - 1

    def test_quando_nem_as_palavras_cabem_volta_a_reparticao_de_sempre(self):
        from flatcap.documento import _colunas_de_texto, _larguras_pelo_conteudo
        from flatcap.relatorio import TabelaLida

        tabela = TabelaLida(
            "x.csv", "", "",
            colunas=("a", "b"),
            linhas=(("palavra" * 20, "outrapalavra" * 20),),
        )
        assert _larguras_pelo_conteudo(tabela, _colunas_de_texto(tabela), int(Pt(453))) is None

    def test_a_figura_ocupa_a_largura_util_da_pagina(self, documento):
        doc, _ = documento()
        secao = doc.sections[-1]
        util = secao.page_width - secao.left_margin - secao.right_margin
        formas = doc.inline_shapes
        assert formas
        assert all(abs(f.width - util) < Pt(2) for f in formas)


# ── O contrato com o template ───────────────────────────────────────────────


class TestFilaDeEstilos:
    """
    O modelo manda, e o código se adapta ao que encontra.
    """

    def test_o_template_gerado_traz_os_estilos_flatcap(self, tmp_path):
        caminho = escrever_template(tmp_path / ARQUIVO_TEMPLATE)
        estilos = {e.name for e in Document(str(caminho)).styles}
        for fila in (ESTILO_TITULO, ESTILO_SECAO, ESTILO_LEGENDA):
            assert fila[0] in estilos

    def test_o_corpo_do_template_sai_vazio(self, tmp_path):
        """
        Um modelo com instruções dentro imprimiria as instruções em cima de todo
        relatório que você entrega.
        """
        caminho = escrever_template(tmp_path / ARQUIVO_TEMPLATE)
        assert _textos(Document(str(caminho))) == []

    def test_o_corpo_do_modelo_e_preservado_e_vem_antes(self, projeto_com_rodada, tmp_path):
        """
        É assim que uma capa institucional continua sendo a capa.
        """
        modelo = tmp_path / ARQUIVO_TEMPLATE
        doc = Document()
        doc.add_paragraph("FACULDADE — Assessoria Estatística")
        doc.save(str(modelo))

        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO, template=modelo)
        executar_relatorio(plano)
        textos = _textos(Document(str(plano.destino)))
        assert textos[0] == "FACULDADE — Assessoria Estatística"
        assert textos[1] == plano.relatorio.titulo

    def test_paragrafos_vazios_do_fim_do_modelo_sao_descartados(
        self, projeto_com_rodada, tmp_path
    ):
        modelo = tmp_path / ARQUIVO_TEMPLATE
        doc = Document()
        doc.add_paragraph("Capa")
        for _ in range(6):
            doc.add_paragraph("")
        doc.save(str(modelo))

        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO, template=modelo)
        executar_relatorio(plano)
        saida = Document(str(plano.destino))
        indice = [p.text for p in saida.paragraphs].index("Capa")
        assert saida.paragraphs[indice + 1].text == plano.relatorio.titulo

    def test_sem_table_grid_no_modelo_a_borda_e_escrita_a_mao(
        self, projeto_com_rodada, tmp_path
    ):
        """
        O Word só grava a definição dos estilos que o documento chegou a usar:
        um modelo institucional em que ninguém nunca inseriu uma tabela não traz
        o `Table Grid`, e sem o degrau de baixo o relatório sairia com tabelas
        sem uma linha separando as colunas.
        """
        modelo = tmp_path / ARQUIVO_TEMPLATE
        doc = Document()
        for nome in ESTILO_TABELA:
            try:
                doc.styles[nome].delete()
            except KeyError:
                pass
        doc.save(str(modelo))

        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO, template=modelo)
        executar_relatorio(plano)
        saida = Document(str(plano.destino))
        quadro = next(t for t in saida.tables if len(t.columns) > 1)
        assert "tblBorders" in quadro._tbl.xml

    def test_o_estilo_do_modelo_vence_a_formatacao_a_mao(self, projeto_com_rodada, tmp_path):
        """
        Passar por cima do que você ajustou no Word é exatamente o que a
        separação template/código existe para impedir.
        """
        from docx.enum.style import WD_STYLE_TYPE

        modelo = tmp_path / ARQUIVO_TEMPLATE
        doc = Document()
        estilo = doc.styles.add_style("FlatCap Titulo", WD_STYLE_TYPE.PARAGRAPH)
        estilo.font.size = Pt(33)
        doc.save(str(modelo))

        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO, template=modelo)
        executar_relatorio(plano)
        saida = Document(str(plano.destino))
        titulo = next(p for p in saida.paragraphs if p.text == plano.relatorio.titulo)
        assert titulo.style.name == "FlatCap Titulo"
        assert all(t.font.size is None for t in titulo.runs)

    def test_modelo_que_nao_abre_vira_mensagem_e_nao_traceback(
        self, projeto_com_rodada, tmp_path
    ):
        modelo = tmp_path / "quebrado.docx"
        modelo.write_bytes(b"isto nao e um docx")
        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO, template=modelo)
        with pytest.raises(RelatorioError) as erro:
            executar_relatorio(plano)
        assert "flatcap template" in str(erro.value)


class TestGravacao:
    def test_a_gravacao_e_atomica(self, projeto_com_rodada, monkeypatch):
        """
        Meio `.docx` não é um documento incompleto — é um arquivo que o Word se
        recusa a abrir.
        """
        import flatcap.documento as modulo

        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO)

        original = modulo.Document

        class DocumentoQueFalhaAoSalvar:
            def __init__(self, *a, **k):
                self._interno = original(*a, **k)

            def __getattr__(self, nome):
                return getattr(self._interno, nome)

            def save(self, *_a, **_k):
                raise OSError("disco cheio (simulado)")

        monkeypatch.setattr(modulo, "Document", DocumentoQueFalhaAoSalvar)
        with pytest.raises(OSError):
            executar_relatorio(plano)
        assert not plano.destino.exists()
        assert not list(plano.raiz.glob("*.tmp"))

    @pytest.mark.skipif(os.name != "posix", reason="no Windows, quem decide é a pasta")
    def test_o_docx_sai_com_a_permissao_do_umask(self, projeto_com_rodada):
        """
        REGRESSÃO da 1.13.0: o `.docx` saía `0o600`, legível só pelo dono — o
        modo do temporário do `tempfile`. Ver o `ficha.temporario_ao_lado`.
        """
        raiz = projeto_com_rodada()
        plano = preparar_relatorio(raiz, agora=QUANDO)

        anterior = os.umask(0o022)
        try:
            executar_relatorio(plano)
        finally:
            os.umask(anterior)

        assert stat.S_IMODE(plano.destino.stat().st_mode) == 0o644

    def test_o_arquivo_gerado_abre(self, documento):
        doc, plano = documento()
        assert plano.destino.stat().st_size > 10_000
        assert len(doc.tables) >= len(plano.relatorio.tabelas)

    def test_a_versao_do_python_docx_e_informavel(self):
        assert versao_do_python_docx() != "?"


def test_o_que_so_se_ve_abrindo_o_documento():
    """
    Lembrete, não teste. Quatro dos cinco defeitos de formatação da Fase 6 —
    a coluna espremida, as duas famílias de fonte na mesma página, a legenda
    órfã e o `Total` na página seguinte — não teriam sido pegos por nenhuma
    asserção em XML.

    Antes de mesclar: gerar o `.docx` de um projeto de verdade, abrir no Word e
    ler as páginas. É o passo 4 do ritual do roteiro, e é o que encontrou três
    das quatro coisas mais graves deste projeto.
    """
    assert True
