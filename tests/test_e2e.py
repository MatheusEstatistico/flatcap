# SPDX-FileCopyrightText: 2026 Matheus Fernandes
# SPDX-License-Identifier: GPL-3.0-or-later
"""
Ponta a ponta — o `flatcap` como subprocesso, contra um REDCap de mentira.

    pytest              # a offline. É a de todo dia.
    pytest -m e2e       # esta. Antes de mesclar uma fase.
    pytest -m ""        # as duas.

Esta camada enxerga o que a suíte offline não pode enxergar: **o `requests`
montando o corpo urlencoded**, o `typer` interpretando a linha de comando, **o
código de saída que o shell recebe** e o `entry point` do `pyproject.toml`.

Não é hipotético: foi esta camada que pegou o `fields[0]=`. A suíte offline
substitui o `requests.post` e por isso nunca vê o corpo de verdade.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

import bancos
import servidor_falso
from conftest import acrescentar_a_ficha, escrever_analises, marcar_numericas
from servidor_falso import ERRO_FIELDS_INDEXADOS

pytestmark = pytest.mark.e2e


def flatcap(*argumentos: str, env: Path, cwd: Path | None = None):
    """
    Roda o executável instalado, como o usuário roda.

    `-m flatcap` e não o script `flatcap` de propósito: os dois passam pelo
    mesmo `app`, e este funciona mesmo com o `.venv` desativado.
    """
    ambiente = dict(os.environ)
    ambiente.pop("REDCAP_API_URL", None)
    ambiente.pop("REDCAP_TOKEN", None)
    ambiente.pop("FLATCAP_ROOT", None)
    return subprocess.run(
        [sys.executable, "-m", "flatcap", *argumentos, "--env", str(env)],
        capture_output=True,
        text=True,
        cwd=str(cwd) if cwd else None,
        env=ambiente,
        timeout=180,
    )


REGISTROS = [
    {
        "record_id": "40", "intro": "1", "matricula1": "00740",
        "prof1": "Ana Cláudia Menezes", "matricula2": "", "aluno2": "",
        "database": "banco.csv",
    },
    {
        "record_id": "42", "intro": "1", "matricula1": "00742",
        "prof1": "José Antônio de Sá", "matricula2": "", "aluno2": "",
        "database": "",
    },
]


@pytest.fixture
def anexo(tmp_path) -> bytes:
    caminho = bancos.csv_do_r(tmp_path / "banco.csv")
    return caminho.read_bytes()


# ── A camada de rede de verdade ─────────────────────────────────────────────


class TestPayload:
    def test_as_listas_vao_por_virgula_e_nao_indexadas(self, tmp_path, anexo):
        """
        REGRESSÃO da Fase 1, e a razão de esta suíte existir. A forma indexada
        `fields[0]` que o API Playground sugere vira `%5B0%5D` no corpo
        urlencoded e chega do outro lado como um `fields` com os nomes grudados.
        """
        with servidor_falso.subir(registros=REGISTROS) as api:
            env = api.env(tmp_path)
            resultado = flatcap("listar", env=env)
        assert resultado.returncode == 0, resultado.stderr

        chamada = api.ultima("record")
        assert "%5B" not in chamada.bruto, ERRO_FIELDS_INDEXADOS
        # Os sete do roteiro e os três da primeira página do relatório.
        assert chamada.lista("fields") == [
            "record_id", "intro", "matricula1", "prof1", "matricula2", "aluno2",
            "fin", "outros1", "obj1", "database",
        ]

    def test_o_servidor_recusa_a_forma_indexada(self, tmp_path):
        """
        A trava do servidor falso. Um retorno acidental àquela forma falha aqui,
        e não em produção.
        """
        import urllib.parse
        import urllib.request

        with servidor_falso.subir(registros=REGISTROS) as api:
            corpo = urllib.parse.urlencode(
                {"token": servidor_falso.TOKEN, "content": "record", "fields[0]": "record_id"}
            ).encode()
            pedido = urllib.request.Request(api.url, data=corpo)
            with pytest.raises(Exception) as erro:
                urllib.request.urlopen(pedido)
            assert "400" in str(erro.value)

    def test_token_recusado_vira_painel_e_codigo_1(self, tmp_path):
        with servidor_falso.subir(registros=REGISTROS) as api:
            env = api.env(tmp_path)
            env.write_text(
                env.read_text(encoding="utf-8").replace(servidor_falso.TOKEN, "errado"),
                encoding="utf-8",
                newline="\n",
            )
            resultado = flatcap("doctor", env=env)
        assert resultado.returncode == 1
        assert "Traceback" not in resultado.stderr

    def test_corpo_nulo_e_barrado_com_o_diagnostico(self, tmp_path):
        with servidor_falso.subir(registros=REGISTROS, corpo_nulo=True) as api:
            resultado = flatcap("doctor", env=api.env(tmp_path))
        assert resultado.returncode == 1
        assert "Traceback" not in resultado.stderr

    def test_o_nome_do_anexo_vem_do_content_type(self, tmp_path, anexo):
        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("base do cliente.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            resultado = flatcap("novo", "40", "--sim", env=env)
        assert resultado.returncode == 0
        assert "base do cliente.csv" in " ".join(resultado.stdout.split())


# ── O fluxo inteiro ─────────────────────────────────────────────────────────


class TestFluxoCompleto:
    def test_do_redcap_ao_docx(self, tmp_path, anexo):
        """
        O critério de pronto do roteiro, executado: `listar`, `novo`,
        `dados adicionar`, `perfil`, `descritiva`, `relatorio`.
        """
        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"

            assert flatcap("listar", env=env).returncode == 0
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0

            raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            assert (raiz / "projeto.yaml").is_file()
            assert list((raiz / "1. Dados originais").glob("00_redcap_*"))

            revisado = tmp_path / "corrigido.csv"
            bancos.csv_do_r(revisado, bancos.com_coluna_constante())
            assert flatcap(
                "dados", "adicionar", "40", str(revisado), "--sim", env=env
            ).returncode == 0
            assert list((raiz / "1. Dados originais").glob("01_revisado_*"))

            assert flatcap("perfil", "40", "--sim", env=env).returncode == 0
            assert (raiz / "perfil.md").is_file()
            assert (raiz / "dicionario.csv").is_file()

            assert flatcap("descritiva", "40", "--sim", env=env).returncode == 0
            pasta = raiz / "3. Análises" / "descritiva"
            assert (pasta / "manifesto.json").is_file()

            assert flatcap("relatorio", "40", "--sim", env=env).returncode == 0
            documentos = list(raiz.glob("Relatorio_00740_*.docx"))
            assert len(documentos) == 1
            assert documentos[0].stat().st_size > 10_000

            historico = (raiz / "historico.md").read_text(encoding="utf-8")
            assert historico.count("- **") == 5

    def test_o_numero_do_docx_e_o_do_csv(self, tmp_path, anexo):
        """
        A decisão de arquitetura nº 3 conferida no produto final: um número que
        aparece no documento é, letra por letra, o do `.csv` ao lado.
        """
        from docx import Document

        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"
            for comando in (
                ("novo", "40", "--sim"),
                ("perfil", "40", "--sim"),
                ("descritiva", "40", "--sim"),
                ("relatorio", "40", "--sim"),
            ):
                assert flatcap(*comando, env=env).returncode == 0

        raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
        pasta = raiz / "3. Análises" / "descritiva"
        linhas = [
            linha
            for linha in (pasta / "freq_sexo.csv").read_text(encoding="utf-8-sig").splitlines()
            if not linha.startswith("#")
        ]
        valores = {linha_.split(";")[0]: linha_.split(";")[2] for linha_ in linhas[1:]}

        documento = Document(str(next(raiz.glob("Relatorio_*.docx"))))
        celulas = {
            linha.cells[0].text.strip(): linha.cells[2].text.strip()
            for tabela in documento.tables
            for linha in tabela.rows
            if len(linha.cells) > 2  # a caixa de aviso tem uma célula só
        }
        for categoria, percentual in valores.items():
            if categoria in celulas:
                assert celulas[categoria] == percentual, categoria


    def test_o_cruzamento_do_docx_e_o_do_csv(self, tmp_path, anexo):
        """
        A decisão de arquitetura nº 3 conferida para a terceira análise: a tabela
        do cruzamento no documento é, célula por célula, a do `.csv` que o
        comando escreveu — e ela entra no documento sem o `relatorio` saber que
        o cruzamento existe.
        """
        from docx import Document

        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            acrescentar_a_ficha(raiz, "cruzamento:\n  - [sexo, escolaridade]\n")
            for comando in (
                ("perfil", "40", "--sim"),
                ("descritiva", "40", "--sim"),
                ("cruzamento", "40", "--sim"),
                ("relatorio", "40", "--sim"),
            ):
                resultado = flatcap(*comando, env=env)
                assert resultado.returncode == 0, resultado.stdout + resultado.stderr

        csv = raiz / "3. Análises" / "cruzamento" / "cruz_sexo_x_escolaridade.csv"
        do_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]

        documento = Document(str(next(raiz.glob("Relatorio_*.docx"))))
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells][:2] == ["sexo", "escolaridade"]
        )
        do_docx = [[c.text.strip() for c in linha.cells] for linha in quadro.rows]
        assert do_docx == do_csv

    def test_a_associacao_do_docx_e_a_do_csv(self, tmp_path, anexo):
        """
        A quarta análise, pela mesma régua: a tabela de testes no documento é,
        célula por célula, a do `.csv` — o p inclusive — e ela entra no documento
        sem o `relatorio` saber que a associação existe.
        """
        from docx import Document

        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            acrescentar_a_ficha(raiz, "associacao:\n  - sexo: [escolaridade]\n")
            for comando in (
                ("perfil", "40", "--sim"),
                ("descritiva", "40", "--sim"),
                ("associacao", "40", "--sim"),
                ("relatorio", "40", "--sim"),
            ):
                resultado = flatcap(*comando, env=env)
                assert resultado.returncode == 0, resultado.stdout + resultado.stderr

        csv = raiz / "3. Análises" / "associacao" / "associacao.csv"
        do_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        assert do_csv[0] == ["variável 1", "variável 2", "n", "teste", "p-valor", "V de Cramér"]

        documento = Document(str(next(raiz.glob("Relatorio_*.docx"))))
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells] == do_csv[0]
        )
        do_docx = [[c.text.strip() for c in linha.cells] for linha in quadro.rows]
        assert do_docx == do_csv

    def test_a_correlacao_do_docx_e_a_do_csv(self, tmp_path, anexo):
        """
        A quinta análise, pela mesma régua: a tabela no documento é, célula por
        célula, a do `.csv` — o coeficiente e o intervalo inclusive —, e a dispersão
        entra como figura, sem o `relatorio` saber que a correlação existe.
        """
        from docx import Document

        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            acrescentar_a_ficha(raiz, "correlacao:\n  - idade: [imc]\n")
            for comando in (
                ("perfil", "40", "--sim"),
                ("descritiva", "40", "--sim"),
                ("correlacao", "40", "--sim"),
                ("relatorio", "40", "--sim"),
            ):
                resultado = flatcap(*comando, env=env)
                assert resultado.returncode == 0, resultado.stdout + resultado.stderr

        csv = raiz / "3. Análises" / "correlacao" / "correlacao.csv"
        do_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        assert do_csv[0] == [
            "variável 1", "variável 2", "n", "método", "coeficiente", "IC 95%", "p-valor",
        ]

        documento = Document(str(next(raiz.glob("Relatorio_*.docx"))))
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells] == do_csv[0]
        )
        do_docx = [[c.text.strip() for c in linha.cells] for linha in quadro.rows]
        assert do_docx == do_csv
        legendas = [p.text for p in documento.paragraphs if p.text.startswith("Figura")]
        assert any("Dispersão — idade × imc" in legenda for legenda in legendas)


    def test_a_comparacao_do_docx_e_a_do_csv(self, tmp_path, anexo):
        """
        A sexta análise, pela mesma régua: a tabela no documento é, célula por célula,
        a do `.csv` — a média, o intervalo e a referência inclusive —, sem o
        `relatorio` saber que a comparação existe.
        """
        from docx import Document

        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            acrescentar_a_ficha(
                raiz,
                "comparacao:\n  - [imc, 25]\n  - [idade, 40, t]\n  - [imc, sexo, student]\n"
                "  - [idade, escolaridade, anova]\n",
            )
            for comando in (
                ("perfil", "40", "--sim"),
                ("descritiva", "40", "--sim"),
                ("comparacao", "40", "--sim"),
                ("relatorio", "40", "--sim"),
            ):
                resultado = flatcap(*comando, env=env)
                assert resultado.returncode == 0, resultado.stdout + resultado.stderr

        csv = raiz / "3. Análises" / "comparacao" / "uma_amostra.csv"
        do_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        assert do_csv[0] == [
            "variável", "n", "média (DP)", "IC 95% da média", "referência", "teste", "p-valor",
        ]
        assert [linha[5] for linha in do_csv[1:]] == ["t de Student", "t de Student"]

        documento = Document(str(next(raiz.glob("Relatorio_*.docx"))))
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells] == do_csv[0]
        )
        do_docx = [[c.text.strip() for c in linha.cells] for linha in quadro.rows]
        assert do_docx == do_csv

        # Entre grupos: a segunda tabela da mesma seção, com o bloco da comparação e
        # uma linha por grupo — as células vazias inclusive.
        csv = raiz / "3. Análises" / "comparacao" / "dois_grupos.csv"
        grupos_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        assert grupos_csv[0] == [
            "variável", "grupo", "n", "média (DP)", "diferença (IC 95%)", "teste", "p-valor",
        ]
        assert [linha[1] for linha in grupos_csv[1:]] == ["sexo", "Feminino", "Masculino"]
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells] == grupos_csv[0]
        )
        assert [[c.text.strip() for c in linha.cells] for linha in quadro.rows] == grupos_csv

        csv = raiz / "3. Análises" / "comparacao" / "k_grupos.csv"
        k_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        assert k_csv[0] == ["variável", "grupo", "n", "média (DP)", "teste", "p-valor"]
        assert [linha[4] for linha in k_csv[1:]] == ["ANOVA", "", "", ""]
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells] == k_csv[0]
        )
        assert [[c.text.strip() for c in linha.cells] for linha in quadro.rows] == k_csv


    def test_o_lote_roda_a_lista_e_o_relatorio_a_segue(self, tmp_path, anexo):
        """
        O `analisar` pelo executável: a lista vazia recusa com o código que o shell
        recebe, a lista escrita roda de uma vez, e o relatório monta as seções na ordem
        dela — o cruzamento antes da descritiva, porque a ficha disse.
        """
        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            raiz = tmp_path / "FlatCapRoot" / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            assert flatcap("perfil", "40", "--sim", env=env).returncode == 0

            vazia = flatcap("analisar", "40", env=env)
            assert vazia.returncode == 1
            assert "analises: [descritiva]" in " ".join(vazia.stderr.replace("│", " ").split())

            escrever_analises(raiz, "analises: [cruzamento, descritiva, comparacao]")
            acrescentar_a_ficha(
                raiz, "cruzamento:\n  - [sexo, escolaridade]\ncomparacao:\n  - [imc, 25]\n"
            )
            lote = flatcap("analisar", "40", "--sem-graficos", "--sim", env=env)
            assert lote.returncode == 0, lote.stdout + lote.stderr
            for pasta in ("cruzamento", "descritiva", "comparacao"):
                assert (raiz / "3. Análises" / pasta / "manifesto.json").is_file()
            historico = (raiz / "historico.md").read_text(encoding="utf-8")
            assert "`flatcap analisar 40` — 3 análises da lista `analises:`" in historico

            relatorio = flatcap("relatorio", "40", "--sim", env=env)
            assert relatorio.returncode == 0, relatorio.stdout + relatorio.stderr
            assert "cruzamento · descritiva · comparacao" in " ".join(relatorio.stdout.split())
            assert list(raiz.glob("Relatorio_00740_*.docx"))

    def test_as_medidas_repetidas_do_docx_e_as_do_csv(self, tmp_path):
        """
        As medidas repetidas pela mesma régua, com um banco que as tem: a tabela no
        documento é, célula por célula, a do `repetidas.csv` — as células vazias das
        linhas de baixo inclusive.
        """
        from docx import Document

        anexo = bancos.csv_do_r(tmp_path / "banco.csv", bancos.para_comparar_repetidas())
        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo.read_bytes())}
        ) as api:
            env = api.env(tmp_path)
            root = tmp_path / "FlatCapRoot"
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            raiz = root / "00740 - Ana Cláudia Menezes" / "Projeto 040"
            acrescentar_a_ficha(
                raiz,
                "comparacao:\n  - repetidas: [pas_0, pas_1, pas_2]\n"
                "  - repetidas: [peso_0, peso_1, peso_2]\n"
                "  - repetidas: [dor_0, dor_1, dor_2]\n",
            )
            resultado = flatcap("perfil", "40", "--sim", env=env)
            assert resultado.returncode == 0, resultado.stdout + resultado.stderr
            # O `perfil` propõe a dor de 0 a 4 como código, e a reunião a marca escala.
            marcar_numericas(raiz, "dor_0", "dor_1", "dor_2")
            for comando in (
                ("descritiva", "40", "--sim"),
                ("comparacao", "40", "--sim"),
                ("relatorio", "40", "--sim"),
            ):
                resultado = flatcap(*comando, env=env)
                assert resultado.returncode == 0, resultado.stdout + resultado.stderr

        csv = raiz / "3. Análises" / "comparacao" / "repetidas.csv"
        do_csv = [
            linha.split(";")
            for linha in csv.read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        # A dor rejeita a normal e vai para o Friedman: a tabela mistura a média e a
        # mediana, e o cabeçalho diz as duas — no `.csv` e, letra por letra, no `.docx`.
        assert do_csv[0] == [
            "medida", "n", "média (DP) ou mediana [Q1 a Q3]", "teste", "p-valor",
        ]
        assert [linha[3] for linha in do_csv[1:]] == [
            "ANOVA de medidas repetidas", "", "",
            "não testado: as medidas não variam dentro dos casos", "", "",
            "Friedman", "", "",
        ]
        assert [linha[2] for linha in do_csv[7:]] == ["1 [0 a 1,25]", "0 [0 a 1]", "0 [0 a 1]"]
        documento = Document(str(next(raiz.glob("Relatorio_*.docx"))))
        quadro = next(
            t for t in documento.tables
            if [c.text.strip() for c in t.rows[0].cells] == do_csv[0]
        )
        assert [[c.text.strip() for c in linha.cells] for linha in quadro.rows] == do_csv


class TestFluxoDaAmostragem:
    """
    O pedido de amostragem de ponta a ponta: o `novo` de um registro sem anexo e,
    depois dele, a amostragem — já com o servidor derrubado, porque do `novo` em
    diante nada fala com o REDCap.
    """

    def test_do_registro_sem_anexo_ao_docx_da_amostra(self, tmp_path):
        from docx import Document

        with servidor_falso.subir(registros=REGISTROS) as api:
            env = api.env(tmp_path)
            assert flatcap("novo", "42", "--sim", env=env).returncode == 0
        env.write_text(
            "REDCAP_API_URL=https://servidor.derrubado.invalido/redcap/api/\n"
            "REDCAP_TOKEN=token_invalido\n"
            f"FLATCAP_ROOT={tmp_path / 'FlatCapRoot'}\n",
            encoding="utf-8",
            newline="\n",
        )

        raiz = tmp_path / "FlatCapRoot" / "00742 - José Antônio de Sá" / "Projeto 042"
        # o `novo` não muda: as três pastas de sempre, e nenhuma quarta
        assert sorted(p.name for p in raiz.iterdir() if p.is_dir()) == [
            "1. Dados originais", "2. Versões antigas", "3. Análises",
        ]

        resultado = flatcap("amostragem", "42", "--sim", env=env)
        assert resultado.returncode == 0, resultado.stdout + resultado.stderr
        pasta = raiz / "4. Amostragem"
        desenho = pasta / "desenho.yaml"
        assert desenho.is_file()

        texto = desenho.read_text(encoding="utf-8")
        desenho.write_text(
            texto.replace("populacao:\n", "populacao: 1200\n"), encoding="utf-8", newline="\n"
        )
        resultado = flatcap("amostragem", "42", "--sim", env=env)
        assert resultado.returncode == 0, resultado.stdout + resultado.stderr
        saida = " ".join(resultado.stdout.split())
        assert "Tamanho da amostra: 292 de 1200" in saida
        # o registro de mentira não tem `fin` nem `obj1`: o documento sai com o aviso
        assert "Leia antes de enviar" in saida

        # o número do documento é, letra por letra, o do `.csv` ao lado; a primeira
        # tabela é a caixa do aviso, e a da conta é a última
        (documento,) = pasta.glob("Amostragem_00742_*.docx")
        quadro = Document(str(documento)).tables[-1]
        do_docx = [[c.text.strip() for c in linha.cells] for linha in quadro.rows]
        do_csv = [
            linha.split(";")
            for linha in (pasta / "tamanho.csv").read_text(encoding="utf-8-sig").splitlines()
            if linha and not linha.startswith("#")
        ]
        assert do_docx == do_csv
        assert do_docx[-1] == ["Tamanho da amostra", "292"]

        # o `desenho.yaml` é o que você escreveu, e o histórico tem as três linhas
        assert desenho.read_text(encoding="utf-8").count("populacao: 1200") == 1
        historico = (raiz / "historico.md").read_text(encoding="utf-8")
        assert historico.count("- **") == 3


class TestSemRede:
    """
    Do `dados adicionar` em diante, nada toca o REDCap — e a prova é rodar com
    o servidor **derrubado** e um token inválido.
    """

    @pytest.fixture
    def projeto_pronto(self, tmp_path, anexo):
        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
        env.write_text(
            "REDCAP_API_URL=https://servidor.derrubado.invalido/redcap/api/\n"
            "REDCAP_TOKEN=token_invalido\n"
            f"FLATCAP_ROOT={tmp_path / 'FlatCapRoot'}\n",
            encoding="utf-8",
            newline="\n",
        )
        raiz = tmp_path / "FlatCapRoot" / "00740 - Ana Cláudia Menezes" / "Projeto 040"
        acrescentar_a_ficha(
            raiz,
            "cruzamento:\n  - [sexo, escolaridade]\n"
            "associacao:\n  - [sexo, escolaridade]\n"
            "correlacao:\n  - [idade, imc]\n"
            "comparacao:\n  - [imc, 25]\n  - [imc, sexo, student]\n",
        )
        return env, raiz

    @pytest.mark.parametrize(
        "comando",
        [
            ("dados", "listar", "40"),
            ("amostragem", "40", "--sim"),
            ("perfil", "40", "--sim"),
            ("descritiva", "40", "--sim"),
            ("normalidade", "40", "--sim"),
            ("cruzamento", "40", "--sim"),
            ("associacao", "40", "--sim"),
            ("correlacao", "40", "--sim"),
            ("comparacao", "40", "--sim"),
            ("relatorio", "40", "--sim"),
            ("template", "--sim"),
        ],
    )
    def test_roda_com_o_servidor_fora_do_ar(self, projeto_pronto, comando):
        env, _ = projeto_pronto
        # os comandos dependem uns dos outros, então rodam em ordem no fixture
        for anterior in (
            ("perfil", "40", "--sim"),
            ("descritiva", "40", "--sim"),
        ):
            if anterior[0] != comando[0]:
                flatcap(*anterior, env=env)
        resultado = flatcap(*comando, env=env)
        assert resultado.returncode == 0, resultado.stdout + resultado.stderr

    def test_o_lote_roda_com_o_servidor_fora_do_ar(self, projeto_pronto):
        """
        As seis análises de uma vez, com o servidor derrubado: o lote é a soma de
        comandos que não tocam a rede, e não pode passar a tocá-la.
        """
        env, raiz = projeto_pronto
        assert flatcap("perfil", "40", "--sim", env=env).returncode == 0
        escrever_analises(
            raiz,
            "analises: [descritiva, normalidade, cruzamento, associacao, correlacao, comparacao]",
        )
        resultado = flatcap("analisar", "40", "--sem-graficos", "--sim", env=env)
        assert resultado.returncode == 0, resultado.stdout + resultado.stderr
        assert len(list((raiz / "3. Análises").glob("*/manifesto.json"))) == 6


class TestCodigosDeSaida:
    """
    O código que o shell recebe. Encadear comandos em script depende dele.
    """

    def test_pasta_ja_existente_sai_com_zero(self, tmp_path, anexo):
        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0
            assert flatcap("novo", "40", "--sim", env=env).returncode == 0

    def test_cadastro_incompleto_sai_com_um(self, tmp_path):
        incompleto = [{"record_id": "50", "intro": "1", "matricula1": "", "prof1": "Sem Nada"}]
        with servidor_falso.subir(registros=incompleto) as api:
            resultado = flatcap("novo", "50", "--sim", env=api.env(tmp_path))
        assert resultado.returncode == 1

    def test_registro_inexistente_sai_com_um(self, tmp_path):
        with servidor_falso.subir(registros=REGISTROS) as api:
            resultado = flatcap("novo", "999", "--sim", env=api.env(tmp_path))
        assert resultado.returncode == 1
        assert "Traceback" not in resultado.stderr


class TestEntryPoint:
    def test_python_m_flatcap_funciona(self):
        """
        `--version` é opção global e não aceita `--env`: ela é `is_eager` e sai
        antes de qualquer configuração ser lida.
        """
        resultado = subprocess.run(
            [sys.executable, "-m", "flatcap", "--version"],
            capture_output=True, text=True, timeout=60,
        )
        assert resultado.returncode == 0
        assert resultado.stdout.strip().startswith("flatcap ")

    def test_o_executavel_declarado_no_pyproject_existe(self, tmp_path):
        import shutil

        assert shutil.which("flatcap"), "o entry point `flatcap` não foi instalado"

    def test_o_relatorio_parte_rapido(self, tmp_path, anexo):
        """
        O efeito colateral de não carregar `pandas` nem `matplotlib` aparece no
        relógio. Um limite folgado, porque a máquina de CI varia — o que se quer
        pegar é a regressão que faz o comando importar meio pacote.
        """
        import time

        with servidor_falso.subir(
            registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
        ) as api:
            env = api.env(tmp_path)
            for comando in (
                ("novo", "40", "--sim"),
                ("perfil", "40", "--sim"),
                ("descritiva", "40", "--sim"),
            ):
                flatcap(*comando, env=env)
            comeco = time.perf_counter()
            resultado = flatcap("relatorio", "40", env=env)
            duracao = time.perf_counter() - comeco

        assert resultado.returncode == 0
        assert duracao < 4.0, f"o relatório levou {duracao:.1f}s — alguém importou algo caro"


def test_o_manifesto_gravado_e_json_valido_e_do_formato_conhecido(tmp_path, anexo):
    with servidor_falso.subir(
        registros=REGISTROS, anexos={"40": ("banco.csv", anexo)}
    ) as api:
        env = api.env(tmp_path)
        for comando in (
            ("novo", "40", "--sim"),
            ("perfil", "40", "--sim"),
            ("descritiva", "40", "--sim"),
        ):
            assert flatcap(*comando, env=env).returncode == 0

    caminho = (
        tmp_path / "FlatCapRoot" / "00740 - Ana Cláudia Menezes" / "Projeto 040"
        / "3. Análises" / "descritiva" / "manifesto.json"
    )
    corpo = json.loads(caminho.read_text(encoding="utf-8"))
    assert corpo["formato"] == 1
    assert corpo["regras"]["semente"] == 42
    assert corpo["ambiente"]["pandas"]
